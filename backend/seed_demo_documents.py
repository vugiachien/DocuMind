from pathlib import Path
from uuid import uuid4

from docx import Document

from app.db.database import SessionLocal
from app.db import models
from app.services.storage_service import storage_service


ROOT = Path(__file__).resolve().parents[1]
PUBLIC_DIR = ROOT / "public"

DEMO_TYPES = [
    {
        "code": "NDA",
        "name": "Bao mat thong tin",
        "description": "Tai lieu bao mat thong tin, NDA, cam ket khong tiet lo.",
        "template": PUBLIC_DIR / "template" / "HAUI_Template_NDA.docx",
        "playbook": PUBLIC_DIR / "playbook" / "HAUI_Playbook_Bao_Mat_Thong_Tin_NDA.docx",
    },
    {
        "code": "SALE",
        "name": "Mua ban hang hoa",
        "description": "Hop dong mua ban hang hoa, giao hang, thanh toan, bao hanh.",
        "template": PUBLIC_DIR / "template" / "HAUI_Template_Mua_Ban_Hang_Hoa.docx",
        "playbook": PUBLIC_DIR / "playbook" / "HAUI_Playbook_Mua_Ban_Hang_Hoa.docx",
    },
]


def _upload_file(path: Path, object_name: str) -> str:
    with path.open("rb") as file_data:
        storage_service.upload_file(
            file_data=file_data,
            length=path.stat().st_size,
            object_name=object_name,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    return object_name


def _extract_playbook_rows(path: Path) -> list[dict]:
    doc = Document(path)
    if not doc.tables:
        return []

    table = doc.tables[0]
    rows = []
    for row in table.rows[1:]:
        cells = [cell.text.strip() for cell in row.cells]
        if len(cells) < 4 or not cells[0]:
            continue
        rows.append(
            {
                "category": cells[0],
                "name": cells[0],
                "description": cells[3],
                "standardClause": cells[1],
                "severity": (cells[2] or "medium").lower(),
                "acceptableDeviation": cells[3],
            }
        )
    return rows


def _upsert_contract_type(db, item: dict) -> models.ContractType:
    contract_type = db.query(models.ContractType).filter(models.ContractType.code == item["code"]).first()
    if not contract_type:
        contract_type = models.ContractType(
            id=str(uuid4()),
            code=item["code"],
            name=item["name"],
            description=item["description"],
        )
        db.add(contract_type)
        db.flush()
    else:
        contract_type.name = item["name"]
        contract_type.description = item["description"]

    template_path = f"templates/{contract_type.id}/{item['template'].name}"
    _upload_file(item["template"], template_path)
    contract_type.templateUrl = template_path
    contract_type.htmlPreview = None
    return contract_type


def _upsert_playbook(db, item: dict, contract_type: models.ContractType) -> models.AuditPolicy:
    playbook_name = item["playbook"].name
    file_path = f"audit_policies/{playbook_name}"
    _upload_file(item["playbook"], file_path)

    playbook = (
        db.query(models.AuditPolicy)
        .filter(models.AuditPolicy.name == playbook_name, models.AuditPolicy.type == "audit_policy")
        .first()
    )
    if not playbook:
        playbook = models.AuditPolicy(
            id=str(uuid4()),
            name=playbook_name,
            fileUrl=file_path,
            status="active",
            agreementTypeId=contract_type.id,
            type="audit_policy",
        )
        db.add(playbook)
        db.flush()
    else:
        playbook.fileUrl = file_path
        playbook.status = "active"
        playbook.agreementTypeId = contract_type.id
        playbook.type = "audit_policy"

    db.query(models.PlaybookRule).filter(models.PlaybookRule.auditPolicyId == playbook.id).delete()
    for rule in _extract_playbook_rows(item["playbook"]):
        db.add(models.PlaybookRule(auditPolicyId=playbook.id, **rule))

    return playbook


def seed_demo_documents():
    db = SessionLocal()
    try:
        seeded = []
        for item in DEMO_TYPES:
            missing = [str(path) for path in (item["template"], item["playbook"]) if not path.exists()]
            if missing:
                raise FileNotFoundError(f"Missing demo document(s): {', '.join(missing)}")

            contract_type = _upsert_contract_type(db, item)
            playbook = _upsert_playbook(db, item, contract_type)
            rule_count = db.query(models.PlaybookRule).filter(models.PlaybookRule.auditPolicyId == playbook.id).count()
            seeded.append((contract_type.code, contract_type.name, playbook.name, rule_count))

        db.commit()
        for code, type_name, playbook_name, rule_count in seeded:
            print(f"Seeded {code}: {type_name} | {playbook_name} | {rule_count} rules")
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    seed_demo_documents()
